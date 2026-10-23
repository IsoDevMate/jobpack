"""Edit CV wording inside an existing .docx without touching fonts, sizes, or layout."""

from copy import deepcopy
from pathlib import Path

from docx import Document  # re-exported for generate_apps.py


def _set_runs(paragraph, pieces):
    """Set run texts in order; extra runs cleared so bold/italic stay on the original runs."""
    runs = paragraph.runs
    if not runs:
        return
    for i, text in enumerate(pieces):
        if i < len(runs):
            runs[i].text = text
    for i in range(len(pieces), len(runs)):
        runs[i].text = ""


def replace_paragraph_startswith(doc, prefix, new_text, bold_first_word=False):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            if bold_first_word and " " in new_text and len(p.runs) >= 2:
                verb, rest = new_text.split(" ", 1)
                _set_runs(p, [verb + " ", rest])
            else:
                _set_runs(p, [new_text])
            return True
    raise KeyError(f"No paragraph starts with: {prefix!r}")


def replace_exact(doc, old, new, bold_first_word=False):
    for p in doc.paragraphs:
        if p.text == old:
            if bold_first_word and " " in new and len(p.runs) >= 2:
                verb, rest = new.split(" ", 1)
                _set_runs(p, [verb + " ", rest])
            else:
                _set_runs(p, [new])
            return True
    raise KeyError(f"No exact paragraph: {old!r}")


def save_copy(src, dest):
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(src)
    return doc, dest
